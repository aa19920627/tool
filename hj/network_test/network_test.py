'''
@Author: 洪建
@Date 2023/6/9 15:28
'''
import os
import sys
import time

from openpyxl.reader.excel import load_workbook
from openpyxl.workbook import Workbook

'''
读取信息，网络测试的代码
'''


class Network_Test:

    def __int__(self):

        # self.file = self.app_path() + "/network_test/data/ip_list.xlsx"
        pass

    def test_network(self):
        '''
        调用windows自带的网络检查
        '''
        ip_list = self.load_ip_list()

        report_list = []
        for i in ip_list:
            '''
            循环测试ip连通性
            '''
            staion_name, mfid, ip = i
            exit_code = os.system('ping %s' % ip)
            if exit_code == 0:
                report_list.append((staion_name, mfid, ip, '正常'))
            else:
                report_list.append((staion_name, mfid, ip, '网路不通'))

        self.generate_report(report_list)  # 将测试结果写入报告

    def app_path(self):

        '''
        判断程序的运行环境，返回对应程序目录
        '''
        if hasattr(sys, "frozen"):  # 判断运行环境为pyinstaller打包后的环境

            return os.path.dirname(sys.executable)  # 使用pyinstaller打包后exe的运行环境

        return os.path.dirname(os.path.dirname(__file__))  # 使用打包前的目录，用于本地运行

    def load_ip_list(self):
        '''
        读取ip表
        '''
        wb = load_workbook(self.app_path() + "/network_test/data/ip_list.xlsx")  # 获取工作簿对象
        ws = wb.get_sheet_by_name(u'Sheet1')  # 获取服务工作簿
        ip_list = []
        for row in ws.values:

            if row[0] != None and row[0] != '监测站名称（非必填）':
                # 获取服务表中需要的信息
                staion_name = row[0]
                mfid = row[1]
                ip = row[2]
                ip_list.append([staion_name, mfid, ip])
        return ip_list

    def generate_report(self, report_list):
        '''
        生成测试报告
        '''
        book = Workbook()
        sheet = book.active
        sheet.cell(row=1, column=1).value = '监测站名称'
        sheet.cell(row=1, column=2).value = 'mfid'
        sheet.cell(row=1, column=3).value = 'ip'
        sheet.cell(row=1, column=4).value = '测试结果'
        sheet.cell(row=1, column=5).value = '测试时间'

        row = 2
        for i in report_list:
            staion_name, mfid, ip, result = i
            time_write = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(time.time()))
            sheet.cell(row=row, column=1).value = staion_name
            sheet.cell(row=row, column=2).value = mfid
            sheet.cell(row=row, column=3).value = ip
            sheet.cell(row=row, column=4).value = result
            sheet.cell(row=row, column=5).value = time_write
            row = row + 1

        time_now = time.strftime("%Y-%m-%d %H时%M分%S秒", time.localtime(time.time()))
        book.save(self.app_path() + "/network_test/data/%s网络测试结果.xlsx" % time_now)


if __name__ == '__main__':
    nt = Network_Test()
    # Network_Test().test_network('192.168.11.119')
    nt.test_network()
