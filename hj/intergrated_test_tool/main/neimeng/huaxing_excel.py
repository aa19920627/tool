# -*- coding: utf-8 -*-
# @Time    : 2021/8/21 11:04
# @Author  : 洪建
# @FileName: huaxing_excel.py
# @Software: PyCharm
import re

from openpyxl import load_workbook, Workbook


#处理华信的平台导出数据，生成可以注册到华日管控的数据

class Manage_Excel_Data:

    def __init__(self):

        #初始化华信的信息表
        self.work_book = load_workbook('../../data/neimeng/导出报告.xlsx')
        self.sheet = self.work_book['导出报告']

        #华信代理服务通用地址
        self.agency_url = 'http://172.16.246.12:80/srrc_nmg/common/ps/passthrough/passthrough11'

        #初始化注册华日管控的信息表
        self.hr_wb = Workbook(write_only=True)
        self.filepath = '../../data/neimeng/华日注册原子服务信息.xlsx'

#读取华信导出的数据
    def read_data(self):

        info_list = []

        row_line = 1

        for row in self.sheet.rows:

            if row_line != 1:

                mf_name = row[0].value
                vendor = row[1].value
                equip_name = row[2].value
                feature_name = row[4].value
                #正则表达式提取设备能力的英文
                feature_name = re.findall(r"\（(.*?)\）",feature_name)[0]
                url = row[5].value
                ps_code = row[6].value
                bs_code = row[7].value
                mfid = row[8].value
                equipid = row[9].value
                hr_agency_url = self.agency_url + '?' + 'BSCode=%s'%bs_code + '&' + 'PSCode=%s'%ps_code
                # print(mf_name,vendor,equip_name,feature_name,ps_code,bs_code,mfid,equipid)
                # print(hr_agency_url)

                info_list.append([mf_name,vendor,equip_name,feature_name,url,
                                  mfid,equipid,ps_code,bs_code,hr_agency_url])

            row_line +=1

        return  info_list

#新建1个excel表格，存储注册华日平台需要的数据
    def write_new_excel(self):

        #删除默认创建的sheet
        # ws = self.hr_wb['Sheet']
        # self.hr_wb.remove(ws)

        #创建sheet
        hr_ws = self.hr_wb.create_sheet('原子服务信息')

        #循环写入行数据
        for i in self.read_data():

            hr_ws.append(i)

        #保存excel
        self.hr_wb.save(self.filepath)




if __name__ == '__main__':

    mxd = Manage_Excel_Data()

    mxd.write_new_excel()

    # a = '动环设备信息查询（E_QueryDeviceInfo）'
    # print(re.findall(r"\（(.*?)\）",a))