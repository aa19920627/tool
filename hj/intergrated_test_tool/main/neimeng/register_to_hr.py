# -*- coding: utf-8 -*-
# @Time    : 2021/8/21 15:09
# @Author  : 洪建
# @FileName: register_to_hr.py
# @Software: PyCharm


#注册监测站，设备，原子服务到华日
import requests as requests
from openpyxl import load_workbook


class Register_To_Hr:

    def __init__(self):

        #华日管控地址
        self.url_base = '172.16.246.43'

        self.session = requests.session()

        #监测站信息
        self.station_file = '../../data/neimeng/监测站信息.xlsx'
        #设备信息
        self.equip_file = '../../data/neimeng/监测设备信息.xlsx'
        #原子服务信息
        self.feature_file = '../../data/neimeng/华日注册原子服务信息.xlsx'
    #读取监测站信息
    def read_station_info(self):

        ws = load_workbook(self.station_file)
        sheet = ws['Sheet1']

        info_list = []
        info_line = 1

        for i in sheet.rows:

            if info_line !=1 :

                mfid,mf_name,longitude,latitude,address,ip = i

                info_list.append((str(mfid.value),mf_name.value,str(longitude.value),
                                  str(latitude.value),address.value,str(ip.value)))

            info_line +=1

        return info_list

    #注册监测站
    def register_station(self):

        url = 'http://%s:8013/facility' % self.url_base

        for i in self.read_station_info():

            mfid,mf_name,longitude,latitude,address,ip = i

            mftype = '0' + mfid[8]
            areacode = mfid[0:6]
            fmskind = '0'+ mfid[9]

            data = {"mfid":mfid,
                    "areacode":areacode,
                    "ascode":"",
                    "mftype":mftype,
                    "mfname":mf_name,
                    "status":"01",
                    "integratedco":"大公博创",
                    "remocontip":ip,
                    "remocontipport":"",
                    "runningdate":"2020-08-06",
                    "repealdate":"",
                    "latitude":latitude,
                    "longitude":longitude,
                    "address":address,
                    "magangle":"",
                    "fmskind":fmskind,
                    "fmaddrtype":"01",
                    "islink":"01",
                    "remark":"",
                    "mcsdtype":"01",
                    "mcstype":"01",
                    "carlicense":"",
                    "enginenum":""}

            # print(data)
            res = self.session.post(url = url ,json=data)
            print(res.json())

    #读取设备信息
    def read_equip_info(self):

        ws = load_workbook(self.equip_file)
        sheet = ws['Sheet1']

        info_list = []
        info_line = 1

        for i in sheet.rows:

            if info_line !=1 :

                mfname,mfid,equip_name,equipid,equip_type = i

                info_list.append((mfname.value,mfid.value,equip_name.value,equipid.value,
                                  equip_type.value))

            info_line +=1

        return info_list

    #注册设备动环设备
    def register_dh_equip(self,equid,equname,mfid,mfname):

        data = {"equid":equid,
                "equname":equname,
                "equtype":"04",
                "equstatus":"01",
                "equmodel":"",
                "equsn":"",
                "tasknum":1,
                "equip":"",
                "equport":"",
                "equimanu":"大公博创",
                "equworkmode":"",
                "startfreq":"",
                "endfreq":"",
                "channelcount":1,
                "dfband":"1",
                "ifband":"1",
                "power":"1",
                "maxband":"1",
                "modulation":"",
                "mfid":mfid,
                "mfname":mfname,
                "anteids":""}

        return data

    #注册监测或测向设备
    def register_equip(self,equid,equname,equtype,mfid,mfname):

        data = {"equid":equid,
                "equname":equname,
                "equtype":equtype,
                "equstatus":"01",
                "equmodel":"",
                "equsn":"",
                "tasknum":1,
                "equip":"",
                "equport":"",
                "equimanu":"",
                "equworkmode":"",
                "startfreq":"",
                "endfreq":"",
                "channelcount":1,
                "dfband":"1",
                "ifband":"1",
                "power":"1",
                "maxband":"1",
                "modulation":"",
                "mfid":mfid,
                "mfname":mfname,
                "anteids":""}

        return data

    #因为设备类型分为动环、监测设备，测向设备，传参有区别，选择需要注册的设备
    def register_equip_choose(self):

        url = 'http://%s:8013/equip' % self.url_base

        #循环取出需要注册的设备信息，判断设备类型
        for i in self.read_equip_info():

            mfname,mfid,equip_name,equipid,equip_type = i

            if equip_type == 4:

                data = self.register_dh_equip(equipid, equip_name, mfid, mfname)
            else:
                equip_type = '0' + str(equip_type)
                data = self.register_equip(equipid, equip_name, equip_type, mfid, mfname)

            #注册设备
            # print(data)
            res = self.session.post(url=url,json=data)
            print(res.json())

    #读取原子服务信息
    def read_feature_info(self):

        ws = load_workbook(self.feature_file)
        sheet = ws['原子服务信息']

        info_list = []

        for i in sheet.rows:

            featurecode = i[3].value
            url = i[4].value
            mfid = i[5].value
            equid = i[6].value
            pscode = i[7].value
            bscode = i[8].value
            proxyUrl = i[9].value

            info_list.append((featurecode,url,mfid,equid,pscode,bscode,proxyUrl))
            # print(featurecode,url,mfid,equid,pscode,bscode,proxyUrl)
        return info_list

    #注册服务能力
    def register_base_equip(self):

        for i in self.read_feature_info():

            featurecode, feature_url, mfid, equid, pscode, bscode, proxyUrl = i

            url = 'http://%s:8013/equip/%s/atomservices' %(self.url_base,equid)

            data = {"mfid":mfid,
                    "equid":equid,
                    "url":feature_url,
                    "proxyUrl":proxyUrl,
                    "bsCode":bscode,
                    "psCode":pscode,
                    "featurecode":featurecode}

            res = self.session.post(url=url,json=data)
            print(res.json())



if __name__ == '__main__':

    rth = Register_To_Hr()
    #注册监测站
    # rth.register_station()
    #注册监测设备
    # rth.register_equip_choose()
    #注册设备能力
    rth.register_base_equip()